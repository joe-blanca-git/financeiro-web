import io
import logging
from datetime import datetime, date
from decimal import Decimal
from typing import Optional, List, Dict, Any

from fastapi import UploadFile, HTTPException
from sqlalchemy.orm import Session
import openpyxl

from app.models.importacao import Importacao
from app.models.previsao_apb import PrevisaoApb
from app.models.previsao_acr import PrevisaoAcr
from app.models.previsao_titulos_liquidados import PrevisaoTitulosLiquidados
from app.models.user import User

logger = logging.getLogger("santamaria")

class PrevisaoCaixaService:
    def __init__(self, db: Session):
        self.db = db

    def _normalizar_header(self, h: Any) -> str:
        if not h:
            return ""
        s = str(h).strip().lower()
        # Remove caracteres especiais comuns
        s = s.replace("|", "").replace("-", "_").replace(" ", "_")
        return s

    def _parse_data(self, val: Any) -> Optional[date]:
        if val is None or val == "":
            return None
        if isinstance(val, datetime):
            return val.date()
        if isinstance(val, date):
            return val
        s = str(val).strip()
        if not s:
            return None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        return None

    def _parse_decimal(self, val: Any) -> Optional[Decimal]:
        if val is None or val == "":
            return None
        if isinstance(val, (int, float, Decimal)):
            return Decimal(str(val))
        s = str(val).strip().replace("R$", "").replace(" ", "")
        if not s:
            return None
        # Formato brasileiro 1.234,56 -> 1234.56
        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            s = s.replace(",", ".")
        try:
            return Decimal(s)
        except Exception:
            return None

    def _parse_str(self, val: Any, max_length: Optional[int] = None) -> Optional[str]:
        if val is None:
            return None
        s = str(val).strip()
        if not s:
            return None
        if max_length and len(s) > max_length:
            return s[:max_length]
        return s

    def _parse_int(self, val: Any) -> Optional[int]:
        if val is None or val == "":
            return None
        if isinstance(val, int):
            return val
        if isinstance(val, float):
            return int(round(val))
        s = str(val).strip()
        if not s:
            return None
        if s.endswith(".0"):
            s = s[:-2]
        try:
            return int(s)
        except Exception:
            return None

    async def importar_planilha_apb(self, file: UploadFile, current_user: Optional[User] = None) -> Dict[str, Any]:
        """
        Processa e importa o arquivo APB.xlsx:
        1. Cria o registro na tabela importacoes (gerando idImportacoes).
        2. Lê a aba DADOS com formulas avaliadas (data_only=True).
        3. Realiza o insert em lote na tabela previsao_apb vinculado à importação.
        """
        filename = file.filename or "APB.xlsx"
        ext = filename.split(".")[-1].lower() if "." in filename else "xlsx"

        if ext not in ["xlsx", "xls"]:
            raise HTTPException(
                status_code=400,
                detail=f"Formato de arquivo inválido (.{ext}). Envie um arquivo Excel (.xlsx ou .xls)."
            )

        try:
            contents = await file.read()
            # data_only=True para ler valores calculados de fórmulas (ex: VLOOKUP em NOME EMITENTE)
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
        except Exception as e:
            logger.exception("Erro ao abrir arquivo Excel APB: %s", str(e))
            raise HTTPException(
                status_code=400,
                detail=f"Não foi possível abrir o arquivo Excel. Verifique se o arquivo não está corrompido: {str(e)}"
            )

        # Localiza a aba DADOS (ou busca case-insensitive)
        sheet_target = None
        for name in wb.sheetnames:
            if name.strip().upper() == "DADOS":
                sheet_target = name
                break
        
        if not sheet_target:
            # Fallback se houver apenas uma aba
            if len(wb.sheetnames) == 1:
                sheet_target = wb.sheetnames[0]
            else:
                raise HTTPException(
                    status_code=400,
                    detail=f"A aba 'DADOS' não foi encontrada na planilha. Abas disponíveis: {', '.join(wb.sheetnames)}"
                )

        sheet = wb[sheet_target]
        rows_iter = sheet.iter_rows(values_only=True)

        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise HTTPException(status_code=400, detail="A planilha está vazia.")

        # Mapeamento dinâmico de cabeçalhos
        header_map: Dict[str, int] = {}
        for idx, col_name in enumerate(header_row):
            norm = self._normalizar_header(col_name)
            if norm:
                header_map[norm] = idx

        # Validação básica de colunas essenciais
        alias_dict = {
            "arquivo": ["arquivo"],
            "data": ["data"],
            "cod_tp_fluxo": ["cod_tp_fluxo"],
            "ent_sai": ["ent_sai"],
            "prev_real": ["prev_real"],
            "modulo": ["modulo"],
            "vlr_cam_2": ["vlr_cam_2"],
            "trans": ["trans"],
            "emp": ["emp"],
            "estab": ["estab"],
            "esp": ["esp"],
            "serie": ["serie"],
            "n_titulo": ["n_titulo", "numero_titulo", "titulo"],
            "parcela": ["parcela"],
            "numero_fornecedor": ["numero_fornecedor", "num_fornecedor", "fornecedor"],
            "nome_abrev": ["nome_abrev", "favorecido"],
            "valor_fluxo": ["valor_fluxo"],
            "valor": ["valor"],
            "nome_emitente": ["nome_emitente", "emitente"]
        }

        col_index: Dict[str, Optional[int]] = {}
        for field, aliases in alias_dict.items():
            col_index[field] = None
            for a in aliases:
                if a in header_map:
                    col_index[field] = header_map[a]
                    break

        user_name = current_user.name if current_user and current_user.name else "Sistema"
        user_id = current_user.iduser if current_user else None

        # 1. Criação do lote na tabela importacoes
        nova_importacao = Importacao(
            nomeArquivo=filename,
            extensaoArquivo=ext,
            tipo="CONTAS A PAGAR (APB)",
            idUserInc=user_id,
            createdAt=datetime.now()
        )
        self.db.add(nova_importacao)
        self.db.flush()  # Gera o idImportacoes

        id_importacao = nova_importacao.idImportacoes
        data_agora = datetime.now()

        # 2. Leitura e preparação dos registros
        registros_para_inserir: List[Dict[str, Any]] = []
        total_valor_somado = Decimal("0.00")

        for row in rows_iter:
            # Pula linhas totalmente vazias
            if not row or all(v is None for v in row):
                continue

            def get_val(f_name: str) -> Any:
                idx = col_index.get(f_name)
                if idx is not None and idx < len(row):
                    return row[idx]
                return None

            val_valor = self._parse_decimal(get_val("valor"))
            if val_valor:
                total_valor_somado += val_valor

            reg = {
                "arquivo": self._parse_str(get_val("arquivo"), 50),
                "data": self._parse_data(get_val("data")),
                "cod_tp_fluxo": self._parse_str(get_val("cod_tp_fluxo"), 20),
                "ent_sai": self._parse_str(get_val("ent_sai"), 10),
                "prev_real": self._parse_str(get_val("prev_real"), 10),
                "modulo": self._parse_str(get_val("modulo"), 20),
                "vlr_cam_2": self._parse_decimal(get_val("vlr_cam_2")),
                "trans": self._parse_str(get_val("trans"), 100),
                "emp": self._parse_str(get_val("emp"), 10),
                "estab": self._parse_str(get_val("estab"), 10),
                "esp": self._parse_str(get_val("esp"), 10),
                "serie": self._parse_str(get_val("serie"), 20),
                "n_titulo": self._parse_str(get_val("n_titulo"), 50),
                "parcela": self._parse_str(get_val("parcela"), 10),
                "numero_fornecedor": self._parse_str(get_val("numero_fornecedor"), 30),
                "nome_abrev": self._parse_str(get_val("nome_abrev"), 100),
                "valor_fluxo": self._parse_decimal(get_val("valor_fluxo")),
                "valor": val_valor,
                "nome_emitente": self._parse_str(get_val("nome_emitente"), 255),
                "idImportacoes": id_importacao,
                "data_criacao": data_agora,
                "usuario_inclusao": user_name
            }
            registros_para_inserir.append(reg)

        if not registros_para_inserir:
            self.db.rollback()
            raise HTTPException(status_code=400, detail="Nenhum registro válido foi encontrado na aba DADOS.")

        # 3. Inserção em lote de alta performance
        try:
            self.db.bulk_insert_mappings(PrevisaoApb, registros_para_inserir)
            self.db.commit()
            self.db.refresh(nova_importacao)
        except Exception as e:
            self.db.rollback()
            logger.exception("Erro ao persistir registros de PrevisaoApb: %s", str(e))
            raise HTTPException(status_code=500, detail=f"Erro ao salvar registros no banco de dados: {str(e)}")

        return {
            "sucesso": True,
            "mensagem": f"Planilha APB importada com sucesso ({len(registros_para_inserir)} registros inseridos).",
            "idImportacoes": id_importacao,
            "nomeArquivo": filename,
            "totalRegistros": len(registros_para_inserir),
            "totalValor": float(total_valor_somado)
        }

    async def importar_planilha_acr(self, file: UploadFile, current_user: Optional[User] = None) -> Dict[str, Any]:
        """
        Processa e importa o arquivo ACR.xlsx:
        1. Cria o registro na tabela importacoes (gerando idImportacoes).
        2. Lê a aba ACR (ou Resumo) com formulas avaliadas (data_only=True).
        3. Realiza o insert em lote na tabela previsao_acr vinculado à importação.
        """
        filename = file.filename or "ACR.xlsx"
        ext = filename.split(".")[-1].lower() if "." in filename else "xlsx"

        if ext not in ["xlsx", "xls"]:
            raise HTTPException(
                status_code=400,
                detail=f"Formato de arquivo inválido (.{ext}). Envie um arquivo Excel (.xlsx ou .xls)."
            )

        try:
            contents = await file.read()
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
        except Exception as e:
            logger.exception("Erro ao abrir arquivo Excel ACR: %s", str(e))
            raise HTTPException(
                status_code=400,
                detail=f"Não foi possível abrir o arquivo Excel. Verifique se o arquivo não está corrompido: {str(e)}"
            )

        # Localiza a aba ACR ou Resumo (ou fallback primeira aba)
        sheet_target = None
        for cand in ["ACR", "RESUMO", "DADOS"]:
            for name in wb.sheetnames:
                if name.strip().upper() == cand:
                    sheet_target = name
                    break
            if sheet_target:
                break

        if not sheet_target:
            sheet_target = wb.sheetnames[0]

        sheet = wb[sheet_target]
        rows_iter = sheet.iter_rows(values_only=True)

        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise HTTPException(status_code=400, detail="A planilha está vazia.")

        header_map: Dict[str, int] = {}
        for idx, col_name in enumerate(header_row):
            norm = self._normalizar_header(col_name)
            if norm:
                header_map[norm] = idx

        alias_dict = {
            "est": ["est", "estab", "estabelecimento"],
            "esp": ["esp", "especie"],
            "ser": ["ser", "serie"],
            "titulo": ["titulo", "título", "t_tulo", "n_titulo", "numero_titulo"],
            "parcela": ["parcela", "_p", "/p", "p"],
            "un_neg": ["un_neg", "unneg", "unidade_negocio"],
            "nr_pedcli": ["nr_pedcli", "nrpedcli", "pedcli", "pedido_cliente"],
            "tipo_pedido": ["tipo_pedido", "tp_pedido"],
            "cliente": ["cliente", "cod_cliente", "codigo_cliente"],
            "nome_abrev": ["nome_abrev", "nome_abreviado"],
            "nome_cliente": ["nome_cliente", "cliente_nome", "razao_social"],
            "grp_cliente": ["grp_cliente", "grupo_cliente"],
            "cliente_matriz": ["cliente_matriz", "matriz"],
            "port": ["port", "portador"],
            "cart": ["cart", "carteira"],
            "repres": ["repres", "representante"],
            "emissao": ["emissao", "emissão", "emiss_o", "dt_emissao"],
            "dt_entrega": ["dt_entrega", "data_entrega"],
            "vencto": ["vencto", "vencimento", "dt_vencto"],
            "dt_agdnto": ["dt_agdnto", "agendamento"],
            "liquidac": ["liquidac", "liquidacao", "dt_liquidacao"],
            "refer": ["refer", "referencia"],
            "tit_banco": ["tit_banco", "titulo_banco"],
            "moeda": ["moeda"],
            "val_original": ["val_original", "valor_original", "vlr_original"],
            "saldo": ["saldo", "vlr_saldo"],
            "saldo_apresentacao": ["saldo_apresentacao", "saldo_apresenta_o", "saldo_apresentação"],
            "cond_pagto": ["cond_pagto", "condicao_pagamento"],
            "vl_descto": ["vl_descto", "desconto", "valor_desconto"],
            "vl_abat": ["vl_abat", "abatimento", "valor_abatimento"],
            "dias": ["dias"],
            "pedido_venda": ["pedido_venda", "ped_venda"],
            "data_indic": ["data_indic", "dt_indic"],
            "cond_cobr": ["cond_cobr", "condicao_cobranca"],
            "cidade": ["cidade"],
            "telefone": ["telefone"],
            "cidade_cobranca_cliente": ["cidade_cobranca_cliente", "cidade_cobranca", "cidade_cobran_a_cliente"],
            "id_federal": ["id_federal", "cnpj", "cpf", "cnpj_cpf"],
            "unid_federac": ["unid_federac", "uf", "estado"]
        }

        col_index: Dict[str, Optional[int]] = {}
        for field, aliases in alias_dict.items():
            col_index[field] = None
            for a in aliases:
                if a in header_map:
                    col_index[field] = header_map[a]
                    break

        user_name = current_user.name if current_user and current_user.name else "Sistema"
        user_id = current_user.iduser if current_user else None

        nova_importacao = Importacao(
            nomeArquivo=filename,
            extensaoArquivo=ext,
            tipo="CONTAS A RECEBER (ACR)",
            idUserInc=user_id,
            createdAt=datetime.now()
        )
        self.db.add(nova_importacao)
        self.db.flush()

        id_importacao = nova_importacao.idImportacoes
        data_agora = datetime.now()

        registros_para_inserir: List[Dict[str, Any]] = []
        total_valor_somado = Decimal("0.00")

        for row in rows_iter:
            if not row or all(v is None for v in row):
                continue

            def get_val(f_name: str) -> Any:
                idx = col_index.get(f_name)
                if idx is not None and idx < len(row):
                    return row[idx]
                return None

            val_saldo = self._parse_decimal(get_val("saldo"))
            val_original = self._parse_decimal(get_val("val_original"))
            if val_saldo:
                total_valor_somado += val_saldo
            elif val_original:
                total_valor_somado += val_original

            reg = {
                "est": self._parse_str(get_val("est"), 20),
                "esp": self._parse_str(get_val("esp"), 20),
                "ser": self._parse_str(get_val("ser"), 20),
                "titulo": self._parse_str(get_val("titulo"), 50),
                "parcela": self._parse_str(get_val("parcela"), 20),
                "un_neg": self._parse_str(get_val("un_neg"), 50),
                "nr_pedcli": self._parse_str(get_val("nr_pedcli"), 50),
                "tipo_pedido": self._parse_str(get_val("tipo_pedido"), 20),
                "cliente": self._parse_str(get_val("cliente"), 50),
                "nome_abrev": self._parse_str(get_val("nome_abrev"), 100),
                "nome_cliente": self._parse_str(get_val("nome_cliente"), 255),
                "grp_cliente": self._parse_str(get_val("grp_cliente"), 50),
                "cliente_matriz": self._parse_str(get_val("cliente_matriz"), 50),
                "port": self._parse_str(get_val("port"), 50),
                "cart": self._parse_str(get_val("cart"), 50),
                "repres": self._parse_str(get_val("repres"), 50),
                "emissao": self._parse_data(get_val("emissao")),
                "dt_entrega": self._parse_data(get_val("dt_entrega")),
                "vencto": self._parse_data(get_val("vencto")),
                "dt_agdnto": self._parse_data(get_val("dt_agdnto")),
                "liquidac": self._parse_data(get_val("liquidac")),
                "refer": self._parse_str(get_val("refer"), 100),
                "tit_banco": self._parse_str(get_val("tit_banco"), 50),
                "moeda": self._parse_str(get_val("moeda"), 20),
                "val_original": val_original,
                "saldo": val_saldo,
                "saldo_apresentacao": self._parse_decimal(get_val("saldo_apresentacao")),
                "cond_pagto": self._parse_str(get_val("cond_pagto"), 50),
                "vl_descto": self._parse_decimal(get_val("vl_descto")),
                "vl_abat": self._parse_decimal(get_val("vl_abat")),
                "dias": self._parse_int(get_val("dias")),
                "pedido_venda": self._parse_str(get_val("pedido_venda"), 50),
                "data_indic": self._parse_data(get_val("data_indic")),
                "cond_cobr": self._parse_str(get_val("cond_cobr"), 50),
                "cidade": self._parse_str(get_val("cidade"), 100),
                "telefone": self._parse_str(get_val("telefone"), 50),
                "cidade_cobranca_cliente": self._parse_str(get_val("cidade_cobranca_cliente"), 100),
                "id_federal": self._parse_str(get_val("id_federal"), 50),
                "unid_federac": self._parse_str(get_val("unid_federac"), 10),
                "idImportacoes": id_importacao,
                "data_criacao": data_agora,
                "usuario_inclusao": user_name
            }
            registros_para_inserir.append(reg)

        if not registros_para_inserir:
            self.db.rollback()
            raise HTTPException(status_code=400, detail="Nenhum registro válido foi encontrado na planilha ACR.")

        try:
            self.db.bulk_insert_mappings(PrevisaoAcr, registros_para_inserir)
            self.db.commit()
            self.db.refresh(nova_importacao)
        except Exception as e:
            self.db.rollback()
            logger.exception("Erro ao persistir registros de PrevisaoAcr: %s", str(e))
            raise HTTPException(status_code=500, detail=f"Erro ao salvar registros no banco de dados: {str(e)}")

        return {
            "sucesso": True,
            "mensagem": f"Planilha ACR importada com sucesso ({len(registros_para_inserir)} registros inseridos).",
            "idImportacoes": id_importacao,
            "nomeArquivo": filename,
            "totalRegistros": len(registros_para_inserir),
            "totalValor": float(total_valor_somado)
        }

    async def importar_planilha_titulos_liquidados(self, file: UploadFile, current_user: Optional[User] = None) -> Dict[str, Any]:
        """
        Processa e importa o arquivo Titulos liquidados.xlsx:
        1. Cria o registro na tabela importacoes (gerando idImportacoes).
        2. Lê a aba liquidados (ou primeira aba disponível) com formulas avaliadas (data_only=True).
        3. Realiza o insert em lote na tabela previsao_titulos_liquidados em batches de 5.000.
        """
        filename = file.filename or "Titulos_Liquidados.xlsx"
        ext = filename.split(".")[-1].lower() if "." in filename else "xlsx"

        if ext not in ["xlsx", "xls"]:
            raise HTTPException(
                status_code=400,
                detail=f"Formato de arquivo inválido (.{ext}). Envie um arquivo Excel (.xlsx ou .xls)."
            )

        try:
            contents = await file.read()
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
        except Exception as e:
            logger.exception("Erro ao abrir arquivo Excel Titulos Liquidados: %s", str(e))
            raise HTTPException(
                status_code=400,
                detail=f"Não foi possível abrir o arquivo Excel. Verifique se o arquivo não está corrompido: {str(e)}"
            )

        sheet_target = None
        for cand in ["LIQUIDADOS", "TITULOS LIQUIDADOS", "RESUMO", "DADOS"]:
            for name in wb.sheetnames:
                if name.strip().upper() == cand:
                    sheet_target = name
                    break
            if sheet_target:
                break

        if not sheet_target:
            sheet_target = wb.sheetnames[0]

        sheet = wb[sheet_target]
        rows_iter = sheet.iter_rows(values_only=True)

        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise HTTPException(status_code=400, detail="A planilha está vazia.")

        user_name = current_user.name if current_user and current_user.name else "Sistema"
        user_id = current_user.iduser if current_user else None

        nova_importacao = Importacao(
            nomeArquivo=filename,
            extensaoArquivo=ext,
            tipo="TITULOS LIQUIDADOS",
            idUserInc=user_id,
            createdAt=datetime.now()
        )
        self.db.add(nova_importacao)
        self.db.flush()

        id_importacao = nova_importacao.idImportacoes
        data_agora = datetime.now()

        registros_batch: List[Dict[str, Any]] = []
        total_registros = 0
        total_valor_somado = Decimal("0.00")
        BATCH_SIZE = 5000

        try:
            for row in rows_iter:
                if not row or all(v is None for v in row):
                    continue

                def col(idx: int) -> Any:
                    return row[idx] if idx < len(row) else None

                val_mov = self._parse_decimal(col(22))
                val_liq = self._parse_decimal(col(26))
                val_orig = self._parse_decimal(col(37))

                if val_mov:
                    total_valor_somado += val_mov
                elif val_liq:
                    total_valor_somado += val_liq
                elif val_orig:
                    total_valor_somado += val_orig

                reg = {
                    "estabelecimento": self._parse_str(col(0), 20),
                    "especie_documento": self._parse_str(col(1), 20),
                    "serie_documento": self._parse_str(col(2), 20),
                    "titulo": self._parse_str(col(3), 50),
                    "parcela": self._parse_str(col(4), 20),
                    "cliente": self._parse_str(col(5), 50),
                    "nome_cliente": self._parse_str(col(6), 255),
                    "portador": self._parse_str(col(7), 50),
                    "matriz": self._parse_str(col(8), 100),
                    "representante": self._parse_str(col(9), 50),
                    "nome_representante": self._parse_str(col(10), 255),
                    "carteira": self._parse_str(col(11), 50),
                    "data_vencimento": self._parse_data(col(12)),
                    "data_credito": self._parse_data(col(13)),
                    "data_emissao": self._parse_data(col(14)),
                    "trans_abrev": self._parse_str(col(15), 50),
                    "unid_negocio": self._parse_str(col(16), 50),
                    "ad": self._parse_str(col(17), 50),
                    "acumula_atraso": self._parse_str(col(18), 20),
                    "nome_abreviado": self._parse_str(col(19), 100),
                    "descricao_especie": self._parse_str(col(20), 100),
                    "referencia": self._parse_str(col(21), 100),
                    "val_movimento": val_mov,
                    "val_antecipado": self._parse_decimal(col(23)),
                    "descto_abatim": self._parse_decimal(col(24)),
                    "juros": self._parse_decimal(col(25)),
                    "valor_liquido": val_liq,
                    "valor_pendente": self._parse_decimal(col(27)),
                    "retencoes": self._parse_decimal(col(28)),
                    "liquidacao": self._parse_data(col(29)),
                    "vl_cheque": self._parse_decimal(col(30)),
                    "grupo_cliente": self._parse_str(col(31), 50),
                    "numero_plaqueta": self._parse_str(col(32), 50),
                    "ult_aprop_desp_fin": self._parse_str(col(33), 100),
                    "data_transacao": self._parse_data(col(34)),
                    "cliente_matriz": self._parse_str(col(36), 50),
                    "valor_original": val_orig,
                    "processo_exportacao": self._parse_str(col(38), 50),
                    "soma_valor_original": self._parse_str(col(39), 20),
                    "pedido_venda": self._parse_str(col(40), 50),
                    "idImportacoes": id_importacao,
                    "data_criacao": data_agora,
                    "usuario_inclusao": user_name
                }
                registros_batch.append(reg)
                total_registros += 1

                if len(registros_batch) >= BATCH_SIZE:
                    self.db.bulk_insert_mappings(PrevisaoTitulosLiquidados, registros_batch)
                    registros_batch.clear()

            if registros_batch:
                self.db.bulk_insert_mappings(PrevisaoTitulosLiquidados, registros_batch)
                registros_batch.clear()

            if total_registros == 0:
                self.db.rollback()
                raise HTTPException(status_code=400, detail="Nenhum registro válido foi encontrado na planilha de Títulos Liquidados.")

            self.db.commit()
            self.db.refresh(nova_importacao)

        except Exception as e:
            self.db.rollback()
            logger.exception("Erro ao persistir registros de PrevisaoTitulosLiquidados: %s", str(e))
            raise HTTPException(status_code=500, detail=f"Erro ao salvar registros no banco de dados: {str(e)}")

        return {
            "sucesso": True,
            "mensagem": f"Planilha de Títulos Liquidados importada com sucesso ({total_registros} registros inseridos).",
            "idImportacoes": id_importacao,
            "nomeArquivo": filename,
            "totalRegistros": total_registros,
            "totalValor": float(total_valor_somado)
        }
